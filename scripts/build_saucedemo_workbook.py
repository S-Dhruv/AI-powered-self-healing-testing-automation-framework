"""One-off generator for SauceDemo_Test_Cases.xlsx, built to match the exact
column set, header styling, and sheet layout of AI_Assignment_Test_Cases.xlsx
(openpyxl only, per claude.md). Not part of the framework itself -- a data
file producer, run once by hand when the workbook needs regenerating.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "SauceDemo_Test_Cases.xlsx"

HEADERS = [
    "TestCase ID",
    "Title",
    "Test Step",
    "Test Data",
    "Expected Result",
    "Execute",
    "Execution Result",
]

HEADER_FILL = PatternFill(start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFFFF", size=11)
COLUMN_WIDTHS = {"A": 13.0, "B": 34.0, "C": 55.0, "D": 38.0, "E": 55.0, "F": 10.0, "G": 16.33}

ROWS = [
    {
        "id": "TC_001",
        "title": "Verify login page loads successfully",
        "steps": "1. Navigate to https://www.saucedemo.com\n2. Observe the page",
        "data": "URL: https://www.saucedemo.com",
        "expected": "Login page loads; 'Swag Labs' logo, username field, password field, and 'Login' button are visible.",
    },
    {
        "id": "TC_002",
        "title": "Login with valid credentials",
        "steps": "1. Enter username and password\n2. Click 'Login'",
        "data": "Username: standard_user\nPassword: secret_sauce",
        "expected": "'Products' heading and the product list are visible.",
    },
    {
        "id": "TC_003",
        "title": "Login with a locked-out user",
        "steps": "1. Enter username and password\n2. Click 'Login'",
        "data": "Username: locked_out_user\nPassword: secret_sauce",
        "expected": "Error 'Epic sadface: Sorry, this user has been locked out.' is displayed; user is not logged in.",
    },
    {
        "id": "TC_004",
        "title": "Login with invalid password",
        "steps": "1. Enter username and password\n2. Click 'Login'",
        "data": "Username: standard_user\nPassword: wrong_password",
        "expected": "Error 'Epic sadface: Username and password do not match any user in this service' is displayed; user is not logged in.",
    },
    {
        "id": "TC_005",
        "title": "Login with blank username and password",
        "steps": "1. Leave Username and Password empty\n2. Click 'Login'",
        "data": "Username: (blank)\nPassword: (blank)",
        "expected": "Error 'Epic sadface: Username is required' is displayed; form is not submitted.",
    },
    {
        "id": "TC_006",
        "title": "Access inventory page directly without logging in",
        "steps": "1. Navigate to https://www.saucedemo.com/inventory.html\n2. Observe the page",
        "data": "URL: https://www.saucedemo.com/inventory.html",
        "expected": "Error 'Epic sadface: You can only access' is displayed; user remains on the login page.",
    },
    {
        "id": "TC_007",
        "title": "Sort products by price, low to high",
        "steps": "1. Login with valid credentials\n2. Select 'Price (low to high)' from the sort dropdown",
        "data": "Username: standard_user\nPassword: secret_sauce",
        "expected": "Products are reordered with the lowest-priced item ('Sauce Labs Onesie', $7.99) shown first.",
    },
    {
        "id": "TC_008",
        "title": "Add a product to the cart",
        "steps": "1. Login with valid credentials\n2. Click 'Add to cart' for the product",
        "data": "Product: Sauce Labs Backpack",
        "expected": "Cart badge shows '1'; the button for that product changes to 'Remove'.",
    },
    {
        "id": "TC_009",
        "title": "Add multiple products to the cart",
        "steps": "1. Login with valid credentials\n2. Click 'Add to cart' for the first product\n3. Click 'Add to cart' for the second product",
        "data": "Product 1: Sauce Labs Backpack\nProduct 2: Sauce Labs Bike Light",
        "expected": "Cart badge shows '2'.",
    },
    {
        "id": "TC_010",
        "title": "Remove a product from the cart",
        "steps": "1. Login with valid credentials\n2. Click 'Add to cart' for the product\n3. Click 'Remove' for the product",
        "data": "Product: Sauce Labs Backpack",
        "expected": "Cart badge is no longer visible; the button for that product changes back to 'Add to cart'.",
    },
    {
        "id": "TC_011",
        "title": "View product detail page",
        "steps": "1. Login with valid credentials\n2. Click the product name for the product",
        "data": "Product: Sauce Labs Backpack",
        "expected": "Product detail page opens showing 'Sauce Labs Backpack', its price '$29.99', and an 'Add to cart' button.",
    },
    {
        "id": "TC_012",
        "title": "Proceed to checkout and continue with shipping info",
        "steps": (
            "1. Login with valid credentials\n"
            "2. Click 'Add to cart' for the product\n"
            "3. Click 'Cart' icon\n"
            "4. Click 'Checkout'\n"
            "5. Enter First Name, Last Name and Zip/Postal Code\n"
            "6. Click 'Continue'"
        ),
        "data": (
            "Product: Sauce Labs Backpack\n"
            "First Name: QA\n"
            "Last Name: Tester\n"
            "Zip/Postal Code: 94107"
        ),
        "expected": "Checkout overview page opens showing 'Sauce Labs Backpack', 'Payment Information', 'Shipping Information', and 'Price Total' sections.",
    },
    {
        "id": "TC_013",
        "title": "Checkout with missing first name",
        "steps": (
            "1. Login with valid credentials\n"
            "2. Click 'Add to cart' for the product\n"
            "3. Click 'Cart' icon\n"
            "4. Click 'Checkout'\n"
            "5. Enter Last Name and Zip/Postal Code\n"
            "6. Click 'Continue'"
        ),
        "data": (
            "Product: Sauce Labs Backpack\n"
            "Last Name: Tester\n"
            "Zip/Postal Code: 94107"
        ),
        "expected": "Error 'Error: First Name is required' is displayed; checkout overview page is not opened.",
    },
    {
        "id": "TC_014",
        "title": "Complete an order end-to-end",
        "steps": (
            "1. Login with valid credentials\n"
            "2. Click 'Add to cart' for the product\n"
            "3. Click 'Cart' icon\n"
            "4. Click 'Checkout'\n"
            "5. Enter First Name, Last Name and Zip/Postal Code\n"
            "6. Click 'Continue'\n"
            "7. Click 'Finish'"
        ),
        "data": (
            "Product: Sauce Labs Backpack\n"
            "First Name: QA\n"
            "Last Name: Tester\n"
            "Zip/Postal Code: 94107"
        ),
        "expected": "'Thank you for your order!' with 'Your order has been dispatched, and will arrive just as fast as the pony can get there!' is shown.",
    },
    {
        "id": "TC_015",
        "title": "Logout a logged-in user",
        "steps": "1. Login with valid credentials\n2. Click 'Open Menu'\n3. Click 'Logout' in the menu",
        "data": "Username: standard_user\nPassword: secret_sauce",
        "expected": "User is logged out; 'Accepted usernames are:' is displayed again.",
    },
]


def build() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 24.0
    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A2"

    wrap = Alignment(wrap_text=True, vertical="top")
    for row_idx, row in enumerate(ROWS, start=2):
        ws.cell(row=row_idx, column=1, value=row["id"])
        ws.cell(row=row_idx, column=2, value=row["title"])
        ws.cell(row=row_idx, column=3, value=row["steps"]).alignment = wrap
        ws.cell(row=row_idx, column=4, value=row["data"]).alignment = wrap
        ws.cell(row=row_idx, column=5, value=row["expected"]).alignment = wrap
        ws.cell(row=row_idx, column=6, value="Y")
        ws.cell(row=row_idx, column=7, value=None)

    wb.create_sheet("Run Summary")

    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} with {len(ROWS)} test cases")


if __name__ == "__main__":
    build()
