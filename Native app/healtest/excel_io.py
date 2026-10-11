"""Excel input (test suites) and output (results written back into the same workbook)."""

import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logger = logging.getLogger(__name__)

STATUS_FILL = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFEB9C", "Skipped": "E2EFDA"}
RESULT_SHEETS = ("Results", "Step Results", "Run Summary")


@dataclass
class TestCase:
    test_case_id: str
    title: str
    test_step: str
    test_data: str
    expected_result: str
    execute: str = "Y"
    row: int = 0


def _canon(header: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


class ExcelSuite:
    def __init__(self, path: str):
        self.path = Path(path)

    def _test_sheet(self, workbook):
        for name in workbook.sheetnames:
            if name not in RESULT_SHEETS and "test" in name.lower() and "case" in name.lower():
                return workbook[name]
        for sheet in workbook.worksheets:
            if sheet.title not in RESULT_SHEETS:
                return sheet
        return workbook.worksheets[0]

    def read(self) -> List[TestCase]:
        if not self.path.exists():
            raise FileNotFoundError(f"Suite not found: {self.path}")
        workbook = load_workbook(self.path)
        sheet = self._test_sheet(workbook)
        header_row, columns = 1, {}
        for r in range(1, min(6, sheet.max_row) + 1):
            canon = [_canon(c.value) for c in sheet[r]]
            if "testcaseid" in canon:
                header_row, columns = r, {name: i + 1 for i, name in enumerate(canon) if name}
                break
        if not columns:
            raise ValueError("No 'Test Case ID' header found in the first rows of the suite")

        def cell(row, *names) -> str:
            for name in names:
                col = columns.get(name)
                if col and sheet.cell(row=row, column=col).value is not None:
                    return str(sheet.cell(row=row, column=col).value)
            return ""

        cases = []
        for r in range(header_row + 1, sheet.max_row + 1):
            case_id = cell(r, "testcaseid").strip()
            if not case_id:
                continue
            cases.append(TestCase(case_id, cell(r, "testcasetitle", "title"), cell(r, "teststep", "teststeps"),
                                  cell(r, "testdata"), cell(r, "expectedresult"),
                                  (cell(r, "execute") or "Y").strip().upper(), r))
        workbook.close()
        logger.info("[EXCEL] read %d test cases from %s (%d marked N)", len(cases), self.path.name,
                    len([c for c in cases if c.execute == "N"]))
        return cases

    # ------------------------------------------------------------------ write back

    def write_results(self, cases: List[TestCase], case_rows: List[Dict[str, Any]], step_rows: List[Dict[str, Any]],
                      run_id: str, summary: Dict[str, Any]) -> Path:
        target = self.path
        try:
            workbook = load_workbook(self.path)
            self._fill(workbook, cases, case_rows, step_rows, run_id, summary)
            workbook.save(self.path)
        except PermissionError:
            target = self.path.with_name(f"{self.path.stem}_results_{run_id}{self.path.suffix}")
            logger.warning("[EXCEL] %s is open in another program; writing results to %s instead", self.path.name, target.name)
            workbook = load_workbook(self.path)
            self._fill(workbook, cases, case_rows, step_rows, run_id, summary)
            workbook.save(target)
        logger.info("[EXCEL] results written to %s", target)
        return target

    def _fill(self, workbook, cases, case_rows, step_rows, run_id, summary):
        border = Border(*(Side(style="thin"),) * 4)
        header_fill = PatternFill("solid", start_color="366092", end_color="366092")

        def sheet(name: str, headers: List[str], widths: List[int], index: Optional[int] = None):
            if name in workbook.sheetnames:
                del workbook[name]
            ws = workbook.create_sheet(name, index) if index is not None else workbook.create_sheet(name)
            for col, (head, width) in enumerate(zip(headers, widths), 1):
                c = ws.cell(row=1, column=col, value=head)
                c.font, c.fill = Font(bold=True, color="FFFFFF"), header_fill
                ws.column_dimensions[c.column_letter].width = width
            return ws

        def paint(ws, row_idx, values, status):
            for col, value in enumerate(values, 1):
                c = ws.cell(row=row_idx, column=col, value=value)
                c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), border
            if status in STATUS_FILL:
                ws.cell(row=row_idx, column=4 if len(values) > 8 else 5).fill = PatternFill(
                    "solid", start_color=STATUS_FILL[status], end_color=STATUS_FILL[status])

        results = sheet("Results", ["Test Case ID", "Title", "Execute", "Status", "Actual Result", "Screenshot Path",
                                    "Execution Time (s)", "Element Source", "Tokens Used", "Run ID", "Executed At"],
                        [12, 34, 9, 11, 60, 28, 12, 22, 11, 16, 20])
        for i, row in enumerate(case_rows, 2):
            paint(results, i, [row["test_case_id"], row["title"], row["execute"], row["status"], row["actual_result"],
                               row["screenshot_path"], row["execution_time"], row["element_source"], row["tokens_used"],
                               run_id, row["executed_at"]], row["status"])

        steps = sheet("Step Results", ["Test Case ID", "Step", "Step Text", "Status", "Actual Result",
                                       "Element Source", "Tokens Used", "Execution Time (s)"],
                      [12, 6, 50, 11, 60, 22, 11, 12])
        for i, row in enumerate(step_rows, 2):
            for col, value in enumerate([row["test_case_id"], row["step_number"], row["step_text"], row["status"],
                                         row["actual_result"], row["element_source"], row["tokens_used"],
                                         row["execution_time"]], 1):
                c = steps.cell(row=i, column=col, value=value)
                c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), border
            if row["status"] in STATUS_FILL:
                steps.cell(row=i, column=4).fill = PatternFill("solid", start_color=STATUS_FILL[row["status"]],
                                                               end_color=STATUS_FILL[row["status"]])

        summary_ws = sheet("Run Summary", ["Metric", "Value"], [34, 28], index=0)
        rows = [("Run ID", run_id), ("Executed At", datetime.now().isoformat(timespec="seconds")),
                ("Suite", self.path.name), ("Test cases in suite", len(cases)),
                ("Pass", summary["Pass"]), ("Fail", summary["Fail"]), ("Blocked", summary["Blocked"]),
                ("Skipped", summary["Skipped"]), ("Total tokens", summary["total_tokens"]),
                ("LLM calls", summary["llm_calls"]), ("Cost (USD)", summary["cost_usd"]),
                ("Cache hit rate (%)", summary["cache_hit_rate"]), ("Elements healed", summary["healed_count"]),
                ("Duration (s)", summary["duration_s"])]
        for i, (k, v) in enumerate(rows, 2):
            summary_ws.cell(row=i, column=1, value=k).font = Font(bold=True)
            summary_ws.cell(row=i, column=2, value=v)

        test_sheet = self._test_sheet(workbook)
        col = next((c.column for c in test_sheet[1] if _canon(c.value) in ("executionresult", "executionstatus")), None)
        if col is None:
            for r in range(1, 6):
                if any(_canon(c.value) == "testcaseid" for c in test_sheet[r]):
                    col = test_sheet.max_column + 1
                    test_sheet.cell(row=r, column=col, value="Execution Result").font = Font(bold=True)
                    break
        if col:
            by_id = {r["test_case_id"]: r["status"] for r in case_rows}
            for case in cases:
                if case.test_case_id in by_id:
                    test_sheet.cell(row=case.row, column=col, value=by_id[case.test_case_id])
