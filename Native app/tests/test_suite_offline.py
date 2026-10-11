"""End-to-end run of the real executor on saucedemo.com with a stand-in LLM (no API key needed).

The stand-in plans steps with simple rules and picks elements by keyword; it exists to test the
framework (actions, login flow, Excel write-back, verification, zero-token replay), not the model.
"""

import re
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

import healtest.executor as executor_module
from healtest.llm import Identification, Plan
from healtest.settings import Settings

ROOT = Path(__file__).resolve().parent.parent


class StandInLLM:
    model, min_confidence = "stand-in", 0.7

    def __init__(self, cfg, masker, usage, limiter=None):
        self.usage = usage

    def _log(self, ctx, agent, reason):
        self.usage.record(ctx, agent, self.model, reason, 100, 10, 1, "success")

    def plan(self, ctx, step, previous, keys, app, page_lines=None, feedback=""):
        s = step.lower()
        self._log(ctx, "step-planner", "first_learning")
        a = []
        if s.startswith("login"):
            a = [{"action": "login"}]
        elif "enter username and password" in s:
            a = [{"action": "type", "target": "username input on the login page", "value": "{{Username}}"},
                 {"action": "type", "target": "password input on the login page", "value": "{{Password}}"}]
        elif s == "click 'login'":
            a = [{"action": "click", "target": "'Login' button on the login page"}]
        elif s.startswith("click 'add to cart'"):
            a = [{"action": "click", "target": "'Add to cart' button for the product named {{Product}}"}]
        elif "shopping cart" in s:
            a = [{"action": "click", "target": "shopping cart link"}]
        elif "open the menu" in s:
            a = [{"action": "click", "target": "menu button"}]
        elif "logout" in s:
            a = [{"action": "click", "target": "'Logout' link in the menu"}]
        elif "'checkout'" in s:
            a = [{"action": "click", "target": "'Checkout' button"}]
        elif "first name" in s:
            a = [{"action": "type", "target": "first name input", "value": "{{First Name}}"},
                 {"action": "type", "target": "last name input", "value": "{{Last Name}}"},
                 {"action": "type", "target": "postal code input", "value": "{{Postal Code}}"}]
        elif "'continue'" in s:
            a = [{"action": "click", "target": "'Continue' button"}]
        return Plan(ok=bool(a), actions=a, input_tokens=100, output_tokens=10, error="no rule")

    def identify(self, ctx, target, action, page, lines, reason, previous="", feedback=""):
        self._log(ctx, "element-identifier", reason)
        words = [w for w in re.findall(r"[a-z]+", target.lower()) if w not in ("input", "button", "link", "the", "on", "page", "for", "named", "in")]
        flat = [re.sub(r"[^a-z]", "", line.lower()) for line in lines]
        t = target.lower()
        kinds = []
        if "button" in t:
            kinds.append(r"\d+\|(button|input:(submit|button)|a\|)")
        if "link" in t:
            kinds.append(r"\d+\|a\|")
        if "radio" in t:
            kinds.append(r"\d+\|(input:radio|label)")
        if "dropdown" in t:
            kinds.append(r"\d+\|select")
        if "textarea" in t:
            kinds.append(r"\d+\|textarea")
        if re.search(r"\binput\b", t) or action == "type":
            kinds.append(r"\d+\|(input:(text|password|email|number|tel|search)|textarea)")
        scores = []
        for f, line in zip(flat, lines):
            score = sum(w in f for w in words)
            if kinds:
                score += 3 if any(re.match(k, line) for k in kinds) else -4
            scores.append(score)
        best = max(range(len(lines)), key=lambda i: scores[i])
        return Identification(ok=True, index=best, confidence=0.95 if scores[best] > 0 else 0.1, input_tokens=100, output_tokens=10)


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(executor_module, "LLMClient", StandInLLM)
    settings = Settings(str(ROOT / "config"))
    settings.framework["element_store"]["location"] = str(tmp_path / "store")
    settings.framework["usage_logging"]["log_dir"] = str(tmp_path / "logs")
    settings.framework["execution"].update(headless=True, screenshot_dir=str(tmp_path / "shots"), action_delay_seconds=0)
    monkeypatch.chdir(tmp_path)
    return settings, tmp_path


def run(settings, suite):
    return executor_module.Executor(settings, "saucedemo", "production").run(str(suite))


def test_integrity_and_zero_token_replay(env):
    settings, tmp = env
    suite = tmp / "integrity.xlsx"
    shutil.copy(ROOT / "suites" / "integrity_checks.xlsx", suite)

    first = run(settings, suite)
    status = {c["test_case_id"]: c["status"] for c in first["cases"]}
    assert status == {"IC_001": "Fail", "IC_002": "Pass", "IC_003": "Fail", "IC_004": "Skipped"}, first["cases"]
    assert first["summary"]["total_tokens"] > 0 and first["summary"]["ai_identified"] > 0

    second = run(settings, suite)
    assert {c["test_case_id"]: c["status"] for c in second["cases"]} == status
    assert second["summary"]["total_tokens"] == 0, "repeat run must not use tokens"
    assert second["summary"]["cache_hit_rate"] == 100.0

    book = load_workbook(suite)
    assert {"Results", "Step Results", "Run Summary"} <= set(book.sheetnames)
    assert [c.value for c in book["Results"]["D"]][1:5] == ["Fail", "Pass", "Fail", "Skipped"]


def test_second_application_suite(env):
    settings, tmp = env
    suite = tmp / "sauce.xlsx"
    shutil.copy(ROOT / "suites" / "saucedemo_suite.xlsx", suite)
    result = run(settings, suite)
    statuses = {c["test_case_id"]: (c["status"], c["actual_result"]) for c in result["cases"]}
    bad = {k: v for k, v in statuses.items() if v[0] != "Pass"}
    assert not bad, bad
