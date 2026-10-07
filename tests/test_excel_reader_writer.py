import openpyxl
import pytest

from framework.excel.reader import ExcelValidationError, read_suite
from framework.excel.writer import RowResult, RunTotals, append_run_summary, write_results


def _make_workbook(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    headers = ["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result", "Execute"]
    ws.append(headers)
    for r in rows:
        ws.append(r)
    # Give column A a distinct fill so we can assert formatting survives the writer.
    from openpyxl.styles import PatternFill
    ws["A1"].fill = PatternFill(start_color="FFFF00", fill_type="solid")
    wb.save(path)


def test_read_suite_parses_rows(tmp_path):
    path = tmp_path / "suite.xlsx"
    _make_workbook(path, [
        ["TC_001", "Home loads", "1. Navigate to https://x.test", "", "Home visible", "Y"],
        ["TC_002", "Skip me", "1. Click 'X'", "", "X visible", "N"],
    ])
    rows = read_suite(path)
    assert len(rows) == 2
    assert rows[0].execute is True
    assert rows[1].execute is False


def test_read_suite_rejects_missing_columns(tmp_path):
    path = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.append(["TestCase ID", "Title"])  # missing required columns
    ws.append(["TC_001", "x"])
    wb.save(path)
    with pytest.raises(ExcelValidationError):
        read_suite(path)


def test_read_suite_rejects_duplicate_ids(tmp_path):
    path = tmp_path / "dup.xlsx"
    _make_workbook(path, [
        ["TC_001", "A", "1. Click 'A'", "", "A visible", "Y"],
        ["TC_001", "B", "1. Click 'B'", "", "B visible", "Y"],
    ])
    with pytest.raises(ExcelValidationError):
        read_suite(path)


def test_write_results_preserves_formatting_and_creates_backup(tmp_path):
    path = tmp_path / "suite.xlsx"
    _make_workbook(path, [["TC_001", "Home loads", "1. Navigate to https://x.test", "", "Home visible", "Y"]])

    result = RowResult(
        row_index=2, status="Pass", actual_result="matched", screenshot_path="",
        execution_time_ms=123, element_source="Cache", tokens_used=0, llm_model=None,
        run_id="run-1", executed_at="2026-10-07T00:00:00Z",
    )
    write_results(path, "Test Cases", [result])

    backups = list(tmp_path.glob("suite.backup-*.xlsx"))
    assert len(backups) == 1

    wb = openpyxl.load_workbook(path)
    ws = wb["Test Cases"]
    assert ws["A1"].fill.start_color.rgb[-6:] == "FFFF00"  # formatting preserved
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    status_col = header.index("Status") + 1
    assert ws.cell(row=2, column=status_col).value == "Pass"


def test_append_run_summary_creates_sheet(tmp_path):
    path = tmp_path / "suite.xlsx"
    _make_workbook(path, [["TC_001", "A", "1. Click 'A'", "", "A visible", "Y"]])

    totals = RunTotals(
        run_id="run-1", suite="suite.xlsx", environment="test",
        start_time="t0", end_time="t1", total=1, passed=1, failed=0, blocked=0, skipped=0,
        total_input_tokens=0, total_output_tokens=0, estimated_cost=0.0,
        cache_hit_rate=1.0, fallback_rate=0.0, ai_lookups=0,
        heals_pending=0, heals_approved=0, heals_rejected=0, duration_ms=10,
    )
    append_run_summary(path, totals)

    wb = openpyxl.load_workbook(path)
    assert "Run Summary" in wb.sheetnames
    values = [c.value for row in wb["Run Summary"].iter_rows() for c in row if c.value]
    assert "run-1" in values
