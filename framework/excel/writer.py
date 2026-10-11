"""Writes results back into the SAME workbook in place with openpyxl, preserving
fonts, fills, borders, widths, merged cells and data validation. Backs up the
original (timestamped copy) before modifying, per spec §5.2. Never pandas.to_excel()."""

from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

from framework.excel.reader import RESULT_COLUMNS

logger = logging.getLogger(__name__)


@dataclass
class RowResult:
    row_index: int
    status: str
    actual_result: str
    screenshot_path: str
    execution_time_ms: int
    element_source: str
    tokens_used: int
    llm_model: str | None
    run_id: str
    executed_at: str


@dataclass
class RunTotals:
    run_id: str
    suite: str
    environment: str
    start_time: str
    end_time: str
    total: int
    passed: int
    failed: int
    blocked: int
    skipped: int
    total_input_tokens: int
    total_output_tokens: int
    estimated_cost: float
    cache_hit_rate: float
    fallback_rate: float
    ai_lookups: int
    heals_pending: int
    heals_approved: int
    heals_rejected: int
    duration_ms: int


def backup_workbook(path: str | Path) -> Path:
    p = Path(path)
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup_path = p.with_name(f"{p.stem}.backup-{ts}{p.suffix}")
    shutil.copy2(p, backup_path)
    return backup_path


def write_results(
    path: str | Path, sheet_name: str, results: list[RowResult]
) -> Path:
    """Writes result columns in place. Backs up first. Only writes to existing
    result columns (creates them at the end of the header row if absent, without
    disturbing any other column's position/formatting)."""
    p = Path(path)
    backup_workbook(p)

    wb = openpyxl.load_workbook(p)
    ws = wb[sheet_name]

    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_index = {name: i + 1 for i, name in enumerate(header_row) if name}  # 1-based

    next_col = len(header_row) + 1
    for col_name in RESULT_COLUMNS:
        if col_name not in col_index:
            ws.cell(row=1, column=next_col, value=col_name)
            col_index[col_name] = next_col
            next_col += 1

    for r in results:
        ws.cell(row=r.row_index, column=col_index["Status"], value=r.status)
        ws.cell(row=r.row_index, column=col_index["Actual Result"], value=r.actual_result)
        ws.cell(row=r.row_index, column=col_index["Screenshot Path"], value=r.screenshot_path)
        ws.cell(row=r.row_index, column=col_index["Execution Time (ms)"], value=r.execution_time_ms)
        ws.cell(row=r.row_index, column=col_index["Element Source"], value=r.element_source)
        ws.cell(row=r.row_index, column=col_index["Tokens Used"], value=r.tokens_used)
        ws.cell(row=r.row_index, column=col_index["LLM Model"], value=r.llm_model)
        ws.cell(row=r.row_index, column=col_index["Run ID"], value=r.run_id)
        ws.cell(row=r.row_index, column=col_index["Executed At"], value=r.executed_at)

    try:
        wb.save(p)
        output_path = p
    except PermissionError:
        # Excel keeps a workbook open exclusively surprisingly often. Never
        # discard a completed run just because write-back is unavailable.
        output_path = p.with_name(
            f"{p.stem}.results-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}{p.suffix}"
        )
        wb.save(output_path)
        logger.warning(
            "Workbook is locked; results were written to recovery copy %s", output_path
        )
    finally:
        wb.close()
    return output_path


def append_run_summary(path: str | Path, totals: RunTotals) -> Path:
    """Appends a 'Run Summary' sheet entry for this run (spec §5.2). Does not
    overwrite prior runs' summaries — each run gets its own labeled block."""
    p = Path(path)
    wb = openpyxl.load_workbook(p)
    sheet_name = "Run Summary"
    if sheet_name not in wb.sheetnames:
        wb.create_sheet(sheet_name)
    ws = wb[sheet_name]

    start_row = ws.max_row + 2 if ws.max_row and ws.cell(row=1, column=1).value else 1
    pass_rate = f"{(totals.passed / totals.total * 100):.1f}%" if totals.total else "0.0%"

    rows = [
        (f"Run Summary: {totals.run_id}", None),
        (None, None),
        ("Execution Summary", None),
        ("Run ID", totals.run_id),
        ("Suite", totals.suite),
        ("Environment", totals.environment),
        ("Start Time", totals.start_time),
        ("End Time", totals.end_time),
        ("Duration (ms)", totals.duration_ms),
        (None, None),
        ("Test Results", None),
        ("Total Test Cases", totals.total),
        ("Passed", totals.passed),
        ("Failed", totals.failed),
        ("Blocked", totals.blocked),
        ("Skipped", totals.skipped),
        ("Pass Rate", pass_rate),
        (None, None),
        ("Token Usage", None),
        ("Total Input Tokens", totals.total_input_tokens),
        ("Total Output Tokens", totals.total_output_tokens),
        ("Estimated Cost (USD)", totals.estimated_cost),
        (None, None),
        ("Resolution", None),
        ("Cache Hit Rate", f"{totals.cache_hit_rate * 100:.1f}%"),
        ("Fallback Rate", f"{totals.fallback_rate * 100:.1f}%"),
        ("AI Lookups", totals.ai_lookups),
        ("Heals Pending", totals.heals_pending),
        ("Heals Approved", totals.heals_approved),
        ("Heals Rejected", totals.heals_rejected),
    ]
    for offset, (label, value) in enumerate(rows):
        ws.cell(row=start_row + offset, column=1, value=label)
        ws.cell(row=start_row + offset, column=2, value=value)

    try:
        wb.save(p)
        output_path = p
    except PermissionError:
        output_path = p.with_name(
            f"{p.stem}.results-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}{p.suffix}"
        )
        wb.save(output_path)
        logger.warning(
            "Workbook is locked; run summary was written to recovery copy %s", output_path
        )
    finally:
        wb.close()
    return output_path
